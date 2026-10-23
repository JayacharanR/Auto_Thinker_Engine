"""Build the requested IEEE A4 project paper from the supplied Word template.

Run with PYTHONPATH=/tmp/auto_thinker_doclibs and the project Python environment.
The normalized template is a format conversion of the user's Strict OOXML file.
"""

from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import numpy as np
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
FIG = OUT / 'figures'
TEMPLATE = OUT / 'template_normalized.docx'
DEST = OUT / 'Auto_Thinker_Engine_Conference_Paper.docx'
SOURCES = [
 'configs/experiment_contract.yaml', 'configs/navigation.yaml',
 'outputs/comparison/carla_right_turn_simple_camera_route/comparison.json',
 'outputs/robustness/robustness.md', 'reports/traffic_curriculum_assessment_20261005.md',
 'reports/overnight_bev_assessment_20261008.md',
 'reports/navigation_m4_repaired_20261020.json',
 'reports/navigation_m4_matched_native_20261020.md',
 'reports/navigation_train_long_drive_20261020.md',
 'reports/navigation_m3_camera_rules_20261020.md',
 'reports/navigation_first_event_study_20261021.md',
 'outputs/navigation_signal_plateau_train_20261022/assessment.json',
 'outputs/navigation_signal_background_full_20261022/assessment.json',
 'outputs/navigation_signal_spatial_full_20261022/assessment.json',
 'outputs/navigation_signal_context_profile_20261022/assessment.json',
 'outputs/navigation_signal_context_fused_profile_20261022/assessment.json',
 'outputs/navigation_signal_context_cuda_profile_20261022/assessment.json',
 'outputs/navigation_signal_positive_cache_20261022/assessment.json',
 'scripts/train_navigation_student.py',
]

plt.rcParams.update({'font.family':'serif','font.serif':['Liberation Serif'],
 'font.size':8,'axes.titlesize':9,'axes.labelsize':8,'xtick.labelsize':7,
 'ytick.labelsize':7,'legend.fontsize':7,'axes.spines.top':False,
 'axes.spines.right':False,'savefig.dpi':300})
BLUE, ORANGE, GREEN, GRAY, PURPLE = '#24577D','#C27726','#39765B','#777777','#72588F'


def save(fig, name):
    fig.savefig(FIG / f'{name}.png',dpi=300,bbox_inches='tight',pad_inches=.06)
    fig.savefig(FIG / f'{name}.svg',bbox_inches='tight',pad_inches=.06)
    plt.close(fig)


def figures():
    FIG.mkdir(exist_ok=True)
    # Project architecture: implemented branches and a visibly marked research arm.
    fig,ax=plt.subplots(figsize=(7.0,2.65));ax.set(xlim=(0,14),ylim=(-.7,5.5));ax.axis('off')
    def box(x,y,w,h,text,color=BLUE,dashed=False):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.06',
          facecolor='#F5F7F9',edgecolor=color,linewidth=.9,linestyle='--' if dashed else '-'))
        ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=8.2)
    def arrow(x1,y1,x2,y2,dashed=False):
        ax.add_patch(FancyArrowPatch((x1,y1),(x2,y2),arrowstyle='-|>',mutation_scale=10,
          color=GRAY,linewidth=.8,linestyle='--' if dashed else '-'))
    ax.text(.05,5.23,'A. Historical DreamerV3 encoder experiment',fontsize=9,color=BLUE)
    box(.1,3.75,2.1,1.0,'Front RGB\n+ route / ego')
    box(2.65,3.75,3.15,1.0,'Visual encoder\nCNN / JEPA / V-JEPA 2')
    box(6.25,3.75,2.1,1.0,'Adapter + RSSM\nlatent model')
    box(8.8,3.75,2.0,1.0,'Imagination\nactor / critic')
    box(11.25,3.75,2.5,1.0,'Discrete actions\nCARLA loop')
    for a,b in [(2.2,2.65),(5.8,6.25),(8.35,8.8),(10.8,11.25)]:arrow(a,4.25,b,4.25)
    ax.text(.05,3.1,'B. Camera-navigation design; perception / planner promotion remains gated',fontsize=9,color=BLUE)
    box(.1,1.7,2.1,1.0,'Three RGB views\n512 × 288 each')
    box(2.65,1.7,3.15,1.0,'Visual encoder +\nsignal features (staged)',ORANGE,True)
    box(6.25,1.7,2.1,1.0,'Recurrent path\n+ speed / stop')
    box(8.8,1.7,2.0,1.0,'Path and speed\ncontroller')
    box(11.25,1.7,2.5,1.0,'Vehicle actuators\nCARLA loop')
    for a,b in [(2.2,2.65),(5.8,6.25),(8.35,8.8),(10.8,11.25)]:arrow(a,2.2,b,2.2)
    box(2.65,.15,3.15,.9,'Calibration + waypoints\nEgo + own-action history',GREEN)
    arrow(4.22,1.08,4.22,1.7)
    box(6.25,.15,4.55,.9,'CV vs. JEPA / RSSM ensembles\n15 plans; learned models rejected',ORANGE,True)
    arrow(8.52,1.08,8.52,1.7,True)
    ax.text(.1,-.43,'Privileged expert / world labels: collection and scoring only; excluded from camera-policy inputs.',fontsize=7.8,color=GRAY)
    save(fig,'architecture')

    fig,ax=plt.subplots(figsize=(3.3,2.1))
    cond=['Clear','Night','Rain','Sunset','Traffic'];x=np.arange(5)
    for k,(label,vals,col) in enumerate([
      ('CNN',[50,0,45,50,20],GRAY),('Custom JEPA',[100,100,95,100,35],BLUE),
      ('V-JEPA 2',[100,100,100,100,20],ORANGE)]):
        ax.plot(x,vals,marker=['s','o','^'][k],label=label,color=col,lw=1.2,ms=3.5)
    ax.set(xticks=x,xticklabels=cond,ylim=(-4,112),ylabel='Episode success (%)')
    ax.legend(loc='lower left',ncol=1,frameon=False);ax.grid(axis='y',alpha=.2)
    fig.tight_layout();save(fig,'robustness')

    d=json.loads((ROOT/'reports/navigation_m4_repaired_20261020.json').read_text())
    enc=[('cnn_scratch','CNN',GRAY),('imagenet_resnet18','ResNet-18',GREEN),
         ('dinov2_small','DINOv2-S',ORANGE),('jepa_comma2k19','Comma JEPA',BLUE),
         ('jepa_carla','CARLA JEPA',PURPLE)]
    print('M4 encoder IDs:',sorted({g['encoder'] for g in d['groups']}))
    # Repository uses these canonical IDs; fail rather than silently drop a series.
    aliases={'imagenet_resnet18':'resnet18','dinov2_small':'dinov2','jepa_comma2k19':'jepa_comma','jepa_carla':'jepa_carla'}
    all_ids={g['encoder'] for g in d['groups']}
    fig,ax=plt.subplots(figsize=(3.3,2.35))
    for key,label,color in enc:
        actual=key if key in all_ids else aliases[key]
        gs=sorted([g for g in d['groups'] if g['encoder']==actual],key=lambda g:g['labels'])
        if len(gs)!=4:raise ValueError(f'Missing curve {actual}')
        xs=[g['labels'] for g in gs];ys=[g['means']['selection_score'] for g in gs]
        ax.plot(xs,ys,'o-',label=label,color=color,ms=3,lw=1.1)
    ax.set(xscale='log',xticks=[549,1516,3857,15712],xticklabels=['549','1,516','3,857','15,712'],
      xlabel='Labeled training timestamps (log scale)',ylabel='Validation selection score (lower better)',ylim=(.34,.63))
    ax.legend(frameon=False,ncol=2,loc='lower left',fontsize=6.8);ax.grid(alpha=.2)
    fig.tight_layout();save(fig,'label_efficiency')

    fig,ax=plt.subplots(figsize=(3.3,2.1))
    names=['CNN','ResNet','DINOv2','CARLA\nJEPA','Comma\nJEPA'];x=np.arange(5)
    ax.bar(x-.16,[0,6,1,3,4],.32,label='Arrivals',color=BLUE)
    ax.bar(x+.16,[0,3,1,0,2],.32,label='Strict successes',color=GREEN)
    ax.set(xticks=x,xticklabels=names,ylim=(0,7),ylabel='Missions out of six')
    ax.legend(frameon=False,ncol=2,loc='upper left');ax.grid(axis='y',alpha=.2)
    fig.tight_layout();save(fig,'native_pilots')

    fig,ax=plt.subplots(figsize=(3.3,2.05))
    labels=['Action only','Constant','Calibrated CV','JEPA ensemble','RSSM ensemble']
    vals=[.1632826,.1728459,.1759686,.1770619,.1772763]
    ax.barh(labels,vals,color=[GREEN,GRAY,ORANGE,BLUE,PURPLE],height=.65)
    ax.invert_yaxis();ax.set(xlim=(0,.205),xlabel='Endpoint Brier score (lower better)')
    for i,v in enumerate(vals):ax.text(v+.002,i,f'{v:.4f}',va='center',fontsize=7)
    ax.grid(axis='x',alpha=.2);fig.tight_layout();save(fig,'world_model')

    fig,(a,b)=plt.subplots(2,1,figsize=(3.3,2.6),sharex=True)
    ns=['Original\nplateau','Extra\nbackgrounds','Spatial\ncontext']
    vals=[1234,1232,1228]
    a.bar(np.arange(3),vals,color=[GRAY,BLUE,GREEN]);a.set(ylim=(1215,1249),ylabel='Correct / 1,248')
    a.axhline(1234,color='#9B3535',ls='--',lw=.8)
    for i,v in enumerate(vals):a.text(i,v+1,str(v),ha='center',fontsize=7)
    bg=[8197,4848,3221];b.bar(np.arange(3),bg,color=[GRAY,BLUE,GREEN])
    b.axhline(3075.8,color='#9B3535',ls='--',lw=.8)
    b.set(xticks=np.arange(3),xticklabels=ns,ylabel='Background detections',ylim=(0,9700))
    for i,v in enumerate(bg):b.text(i,v+180,f'{v:,}',ha='center',fontsize=7)
    a.grid(axis='y',alpha=.15);b.grid(axis='y',alpha=.15)
    fig.tight_layout(h_pad=.5);save(fig,'perception')

    fig,ax=plt.subplots(figsize=(3.3,1.95))
    names=['Shared averages','Compiled CPU','CUDA + GPU trees'];vals=[343.352,49.088,22.869]
    ax.barh(names,vals,color=[GRAY,BLUE,GREEN],height=.6)
    ax.invert_yaxis();ax.set(xlim=(0,405),xlabel='Median feature + head time (ms / image)')
    for i,v in enumerate(vals):ax.text(v+4,i,f'{v:.1f}',va='center',fontsize=7.5)
    ax.grid(axis='x',alpha=.2);fig.tight_layout();save(fig,'runtime')

    # Use the project's recorded integration route, not an imagined driving example.
    image=Image.open(ROOT/'reports/navigation_train_long_drive_20261020.png')
    image.save(FIG/'long_drive.png')


DOC=None

def set_cols(section,n):
    cols=section._sectPr.find(qn('w:cols'))
    if cols is None:cols=OxmlElement('w:cols');section._sectPr.append(cols)
    cols.set(qn('w:num'),str(n));cols.set(qn('w:space'),'360' if n==2 else '720')


def section(n=2,kind=WD_SECTION_START.NEW_PAGE):
    s=DOC.add_section(kind)
    s.page_width=Pt(595.3);s.page_height=Pt(841.9)
    s.top_margin=Pt(54);s.bottom_margin=Pt(72)
    s.left_margin=Pt(44.65);s.right_margin=Pt(44.65)
    s.header_distance=Pt(36);s.footer_distance=Pt(36)
    s.different_first_page_header_footer=False
    set_cols(s,n)
    return s


def p(text='',style='Body Text',size=None,italic=False,bold=False,align=None):
    a=DOC.add_paragraph(style=style)
    r=a.add_run(text);r.bold=bold;r.italic=italic
    if size:r.font.size=Pt(size)
    if align is not None:a.alignment=align
    return a


def h(text):return p(text,'Heading 1')
def sub(text):return p(text,'Heading 2')
def col():DOC.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)


def fig(name,caption,width=3.38):
    a=p('',align=WD_ALIGN_PARAGRAPH.CENTER);a.paragraph_format.space_after=Pt(0)
    a.paragraph_format.keep_with_next=True
    a.add_run().add_picture(str(FIG/f'{name}.png'),width=Inches(width))
    c=p(caption,'figure caption');c.paragraph_format.space_after=Pt(6)
    c.paragraph_format.keep_together=True


def table(title,headers,rows,widths=None,size=8):
    a=p(title,'table head');a.paragraph_format.space_before=Pt(6)
    a.paragraph_format.space_after=Pt(4);a.paragraph_format.keep_with_next=True
    t=DOC.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths is None:widths=[3.38/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for row,texts in [(t.rows[0],headers),*[(t.add_row(),r) for r in rows]]:
        pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
        for c,w,text in zip(row.cells,widths,texts):
            c.width=Inches(w);a=c.paragraphs[0];a.style=DOC.styles['table copy']
            a.alignment=WD_ALIGN_PARAGRAPH.LEFT
            a.paragraph_format.space_after=Pt(2);a.paragraph_format.space_before=Pt(2)
            a.paragraph_format.first_line_indent=Pt(0);a.paragraph_format.line_spacing=1.0
            r=a.add_run(str(text));r.font.size=Pt(size);r.bold=False
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'909090');borders.append(el)
            tcPr.append(borders)
    for cell in t.rows[0].cells:
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EFF2F5');cell._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t


def build():
    global DOC
    DOC=Document(TEMPLATE)
    body=DOC._element.body
    keep=deepcopy(body.sectPr)
    for e in list(body):body.remove(e)
    body.append(keep)
    # Only unfilled personal/acknowledgement fields are bold throughout the file.
    for s in DOC.styles:
        if hasattr(s, 'font'):s.font.bold=False
        if s.type in (1,2):s.font.name='Times New Roman'
    normal=DOC.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
    for s in DOC.styles:
        if s.type==1:
            fmt=s.paragraph_format
            fmt.widow_control=True
            # Captions/references and numbered headings use explicit numbering in this draft.
            num=s.element.find(qn('w:pPr'))
            if num is not None:
                for e in list(num):
                    if e.tag==qn('w:numPr'):num.remove(e)
    DOC.styles['Body Text'].paragraph_format.space_after=Pt(5)
    DOC.styles['Body Text'].paragraph_format.line_spacing=1.0
    DOC.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for sec in DOC.sections:
        for part in [sec.header,sec.footer,sec.first_page_header,sec.first_page_footer]:
            for child in list(part._element):part._element.remove(child)
    s=DOC.sections[0]
    s.page_width=Pt(595.3);s.page_height=Pt(841.9);s.top_margin=Pt(27);s.bottom_margin=Pt(72)
    s.left_margin=Pt(44.65);s.right_margin=Pt(44.65);set_cols(s,1)
    s.different_first_page_header_footer=False
    DOC.core_properties.title='Evaluating Self-Supervised Video Encoders for World-Model Driving in CARLA'
    DOC.core_properties.subject='Auto Thinker Engine research; verified project evidence through 23 October 2026'
    DOC.core_properties.author=''
    DOC.core_properties.last_modified_by=''
    DOC.core_properties.keywords='CARLA, DreamerV3, JEPA, video pretraining, autonomous driving'
    DOC.core_properties.comments='Author and affiliation placeholders are intentionally bold. Technical results are filled from verified local artifacts.'

    p(DOC.core_properties.title,'paper title')
    a=p('[Author name(s)]','Author',bold=True);a.paragraph_format.space_before=Pt(8)
    for field in ['[Department / programme]','[Institution / organization]','[City, country]','[Email address(es) / ORCID]']:
        a=p(field,'Affiliation',bold=True,size=10);a.paragraph_format.space_after=Pt(0)
    section(2,WD_SECTION_START.CONTINUOUS)
    p('Abstract—This paper studies whether self-supervised video representations improve simulated driving under a single laptop GPU budget. Auto Thinker Engine combines DreamerV3 with CARLA through CarDreamer and compares a trainable convolutional encoder with custom dash-camera JEPA and frozen V-JEPA 2 features. Across two right-turn training seeds, the video encoders reach full evaluated success while the convolutional baseline averages fifty percent; weather robustness improves, but traffic performance remains weak. We extend the investigation to synchronized three-camera navigation, five-encoder label-efficiency comparisons, rule counterfactuals and action-conditioned hazard prediction. The corrected study completes sixty-seven fits, yet better offline scores do not predict safer closed-loop behavior. A camera policy completes a familiar 1.376 km route without recorded contacts, while separate traffic-rule tests and learned hazard models fail their qualification gates. Full-distribution signal audits expose a recall versus false-detection tradeoff. Exact preprocessing acceleration reduces a measured feature-and-classifier path to about twenty-three milliseconds per image. The results support representation benefits on constrained tasks while identifying perception, calibration and evaluation gaps that prevent a claim of reliable urban autonomy.','Abstract')
    p('Keywords—autonomous driving, self-supervised learning, joint-embedding predictive architecture, world models, model-based reinforcement learning, CARLA','Keywords')
    h('I. Introduction')
    p('Learning to drive from camera observations requires scene understanding, temporal memory and decisions about when to stop, yield or proceed. Model-based reinforcement learning offers a way to practice decisions in a learned latent model rather than relying only on expensive simulator interaction. DreamerV3 provides this foundation [1], and CARLA and CarDreamer provide controllable driving environments [2], [3].')
    p('The central question is whether real-video pretraining improves learning efficiency and robustness compared with a visual encoder trained from scratch. A second question is whether these gains survive the transition from one maneuver to destination-directed driving with traffic rules. Project measurements come from versioned contracts and audited artifacts [11]. We report both positive and negative results. The contribution is a laptop-scale experimental pipeline with explicit input boundaries, corrected data splits, full-denominator evaluation and reproducible failure analysis, rather than a claim of solved self-driving.')
    col()
    h('II. Related Work and Research Gap')
    p('Table I positions the project against representative primary work. Dreamer learns control through latent imagination [1]. V-JEPA learns predictive video representations without pixel reconstruction [4], while V-JEPA 2 adds action-conditioned planning demonstrations in robotics [5]. These findings motivate, but do not establish, an autonomous-driving benefit. DINOv2 and supervised ResNet provide distinct visual baselines [6], [7].')
    table('TABLE I. Literature review and relevance',
      ['Work','Main contribution','Gap addressed here'],[
      ['DreamerV3 [1]','General latent world-model RL.','Driving rules and visual transfer require direct tests.'],
      ['CARLA / CarDreamer [2], [3]','Driving simulation and world-model task infrastructure.','A successful short task is not long-route reliability.'],
      ['V-JEPA / V-JEPA 2 [4], [5]','Predictive visual pretraining; action-conditioned planning.','Generic video and robot results do not prove traffic understanding.'],
      ['DINOv2 / ResNet [6], [7]','Strong self-supervised and supervised image features.','Needed to test whether the gain is specific to JEPA.'],
      ['Think2Drive [8]','World-model RL for complex CARLA scenarios.','Privileged observations and larger compute differ from camera-only laptop control.'],
      ['Bench2Drive [9]','Ability-specific closed-loop driving benchmark.','Encourages encounter coverage, not only an aggregate route score.'],
      ],[.76,1.12,1.50],7.6)
    p('The unresolved research gap is the connection between representation quality, decision-critical perception and closed-loop reliability. A pretrained encoder can reduce an offline loss while missing a distant traffic light or producing an unsafe action. This project therefore evaluates representation learning, perception, hazard prediction and driving as separate evidence requirements.')

    # Page 2: wide figure plus two-column methodology.
    section(1)
    h('III. Proposed Design and Methodology')
    fig('architecture','Fig. 1. Project architecture. Dashed boxes mark staged signal perception and the predictive-planning branch, whose learned models remain rejected. Solid paths summarize implemented experimental branches. Privileged labels never become camera-policy inputs.',width=7.0)
    section(2,WD_SECTION_START.CONTINUOUS)
    sub('A. Historical Dreamer encoder experiment')
    p('The first branch in Fig. 1 uses DreamerV3 in PyTorch, with CarDreamer supplying observations, rewards and terminal events. A 64 × 64 RGB decoder target and an identical 23-dimensional route/ego vector are retained across encoder arms. The scratch CNN learns jointly with the recurrent state-space model (RSSM). Frozen video encoders supply features through a trainable adapter with a 1,024-dimensional target. An actor and critic learn from imagined latent rollouts; actions use the same discrete actuator table.')
    p('The custom encoder is a ViT-S/16 JEPA pretrained on real dash-camera video, initially using 128-pixel, eight-frame clips. The transferred checkpoint is frozen during Dreamer learning. The V-JEPA 2 arm uses the frozen facebook/vjepa2-vitl-fpc64-256 checkpoint. The historical experiment is a feature-transfer test; it does not deploy the original V-JEPA 2 action-conditioned robot planner. The shared reward includes the previously established time penalty of 0.1, and the comparison does not retrospectively change that reward.')
    sub('B. Multi-camera destination navigation')
    p('The second branch extends the task to three calibrated RGB cameras at yaw angles of 0 and ±60 degrees. Each image is 512 × 288 pixels with a 100-degree field of view. A frame-ID barrier aligns all views with the observation timestamp. Camera calibration, navigation waypoints, ego motion and timestamps form the declared policy inputs. An approved versioned extension adds only the ego vehicle’s previously applied steering, throttle and brake, with frame identifiers; it is identical for CV, RSSM and JEPA comparisons.')
    col()
    sub('C. Path, speed and temporal decisions')
    p('Navigation geometry supplies an ego-frame route centerline. A recurrent visual policy predicts a path correction separately from the speed/stop profile, reducing the burden of learning route geometry from scratch. The physical controller converts the resulting plan to steering, throttle and brake. This separation focuses learning on stopping, yielding and release decisions. It is an engineering design choice, not evidence that the policy has learned traffic rules.')
    p('The nominal runtime uses a 20 Hz physics and sensor cycle, 10 Hz history sampling and 5 Hz planning. These are configured cadences; measured full-system throughput is reported separately. A four-second planning horizon contains eight path/speed waypoints. Predictive-planning experiments compare a grid of three paths and five speed profiles, producing fifteen candidate plans.')
    sub('D. Predictive models and privilege separation')
    p('Constant-velocity forecasts form the explicit planning baseline. Three-member JEPA and RSSM ensembles predict first-contact hazards from past RGB-derived observations, spatial features, ego/navigation inputs, own-action history and prospective commands. Future video is a training target only. Privileged actor states, simulator traffic-light states and expert controls are restricted to collection, supervision and scoring. An unexecuted alternative action has an unknown outcome unless it is physically branched and observed.')
    p('Expert demonstrations, student behavior and bounded action perturbations address covariate shift. Safety interventions censor subsequent outcomes rather than converting an unsafe action into a successful training example. Exact prefix comparisons test replay agreement, while native recorder and command audits check that intended actions were actually applied. Emergency and engineered rule overrides are disabled in the rule-understanding tests.')

    # Page 3: dataset and evaluation.
    section(2)
    h('IV. Datasets and Experimental Setup')
    sub('A. Real video and simulator data')
    p('comma2k19 contains over 33 hours of California highway driving [10]. This project uses a limited chunk for custom JEPA pretraining, not the entire public dataset. Highway appearance, traffic-light visibility and junction topology differ from urban CARLA scenes; no real-world driving transfer is claimed. V-JEPA 2, DINOv2 and ImageNet ResNet inherit different external data and supervision histories, which are reported rather than treated as equal pretraining compute.')
    p('CARLA 0.9.15 provides generated RGB video, vehicle dynamics, configurable traffic and privileged labels for supervision/evaluation. Town01–04 and Town10HD form the training pool, while Town05 is reserved for final driving. Actual traversed junction geometry, rather than only a route’s nominal town or start ID, determines eligibility. Historical split-contaminated results remain archived and are excluded from the corrected label-efficiency claim.')
    table('TABLE II. Main datasets and retained evaluation sets',
      ['Subset','Size / role'],[
      ['Corrected camera imitation','55 TRAIN episodes; 15,712 labels.'],
      ['Offline camera validation','13 episodes; 2,822 timestamps.'],
      ['Nested label budgets','549 / 1,516 / 3,857 / 15,712 timestamps.'],
      ['Signal corpus','2,844 TRAIN images; 21 signal locations.'],
      ['Signal component validation','144 images; 72 housing labels; four reserved locations.'],
      ['Action/hazard recordings','68 verified of 73 attempts: 50 TRAIN, 18 calibration; five unknown failures.'],
      ['Development encounters','300 frozen encounters across six ability families.'],
      ['Town05 final suite','12 routes × 3 traffic seeds = 36 planned missions; final driving pending.'],
      ],[1.35,2.03],8)
    sub('B. Fair comparisons and hardware')
    p('The corrected visual study completes 67 registered fits: five encoders, four nested label fractions and three seeds, plus controls. Nominal budgets of 5%, 10%, 25% and 100% correspond to actual timestamp fractions of 3.49%, 9.65%, 24.55% and 100% after episode-level grouping. Frozen encoders are cached; the trainable CNN continues to receive images. Downstream update budgets match, while pretraining histories, parameter counts and measured wall time can differ.')
    p('Execution uses one NVIDIA RTX 5070 Ti laptop GPU, Python 3.10 and PyTorch 2.11 with CUDA 12.8. CARLA workloads, data collection, model fitting and profiling are resource-bounded. The approved local plan preserves recordings and scientific artifacts, with a cumulative storage limit and a 10 GB minimum-free-space guard.')
    sub('C. Statistical scope')
    p('Two training seeds and six-case pilots support descriptive comparisons, not population-safety claims. Three-seed ranges are not confidence intervals. Even a perfect 36-mission run permits about an 8% one-sided upper failure bound by the rule of three, before accounting for dependence. Location-clustered intervals and paired tests require sufficient predeclared coverage.')
    col()
    h('V. Evaluation Metrics and Integrity')
    sub('A. Closed-loop driving and rule coverage')
    p('The original task reports destination-reached success, route completion, episode return, collisions, lane departures, training steps to a success threshold, wall time and GPU memory. Historical success and collision flags can overlap in traffic; both flags are retained. Later navigation reports distinguish arrival from rule-safe and strict outcomes, using the registered assessor’s complete criteria rather than reinterpreting earlier scores.')
    p('All recorder contact starts remain in headline contact counts. Fault attribution is diagnostic only. Red-light and stop-sign violations are normalized by actual opportunities; zero violations on a route containing zero opportunities provide no rule-handling evidence. Other measurements include blocked time, interventions, infractions per kilometre, route deviation, comfort and latency. Driving Score is supplementary, not a replacement for physical outcomes.')
    sub('B. Offline policy and hazard metrics')
    p('The corrected policy selector combines controller stop recall R, false-hold fraction F and go-speed mean absolute error E in m/s:')
    p('S = 0.5(1 − R + F) + 0.02E.                    (1)',size=9,align=WD_ALIGN_PARAGRAPH.CENTER)
    p('A lower S is preferred. Path-correction error and speed error are also reported. The score is a development selector, not a probability of safe driving. Stop recall alone is insufficient because an agent that always waits can obtain high recall while blocking every route.')
    p('For a predicted event probability p and observed contact label y, the endpoint Brier score is:')
    p('BS = (1/N) Σᵢ (pᵢ − yᵢ)².                         (2)',size=9,align=WD_ALIGN_PARAGRAPH.CENTER)
    p('First-event sequence negative log likelihood, missed contacts and false alarms complement Brier score. Unknown infrastructure failures and unexecuted action outcomes are not relabeled safe. Ensemble disagreement is not automatically calibrated uncertainty.')
    sub('C. Signal perception and inference cost')
    p('Signal evaluation uses score-ordered one-to-one housing matching, summed foreground probability at least 0.5 and phase argmax. A correct result requires both a point inside the audited housing and the correct red/yellow/green phase. Known-background and unknown-region detections are counted separately. The full TRAIN comparison scores 12,362,313 frozen coordinates rather than only a convenient sampled subset. The fixed gate requires at least 1,234/1,248 correct, every original town and phase protected, and at most 3,075.8 known-background detections.')
    p('Latency reports identify exactly what is timed. Feature extraction plus classifier time excludes camera decode, proposal generation, policy control and simulator work unless stated. Profiling speed is therefore not a substitute for a full closed-loop cadence test.')

    # Page 4: legacy and label efficiency.
    section(2)
    h('VI. Results and Analysis')
    sub('A. Right-turn learning and weather robustness')
    p('The historical three-arm right-turn comparison uses seeds 42 and 123. Both frozen-video arms reach 100% evaluated success, whereas the CNN mean is 50% because its seed outcomes are 100% and 0%. Mean route completion is 100% for both JEPA arms versus 78.8% for CNN. Returns are 337.2, 327.8 and 202.0 for custom JEPA, V-JEPA 2 and CNN respectively. These are selected-checkpoint outcomes on a constrained maneuver.')
    fig('robustness','Fig. 2. Zero-shot success under weather and traffic changes, averaged over two training seeds. Traffic remains difficult for every encoder; weather robustness does not imply interaction safety.')
    table('TABLE III. Historical right-turn comparison',
      ['Arm','Success','80% threshold'],[
      ['CNN scratch','50%','32.5k in one seed; other seed never reaches it'],
      ['Custom JEPA','100%','30.0k ± 2.5k steps'],
      ['V-JEPA 2','100%','21.7k ± 9.2k steps'],
      ],[1.0,.58,1.80],8)
    p('Threshold statistics are descriptive across two seeds; the missing CNN threshold is not imputed. One V-JEPA 2 run is partial at about 66.7k environment steps, so this is not a fully matched-budget convergence comparison. The deferred third seed limits robustness conclusions.')
    p('Zero-shot traffic success is only 20%, 35% and 20% for CNN, custom JEPA and V-JEPA 2. Corresponding collision rates are 55%, 65% and 80%. Pretraining improves appearance robustness in these tests, but does not establish correct gap acceptance or safe yielding.')
    col()
    sub('B. Traffic adaptation and privileged BEV results')
    p('A 30k-step medium-first curriculum yields selected medium-task success of 11/30 for CNN and 4/30 for each video arm; all fail the traffic promotion gate. A separate privileged bird’s-eye-view (BEV) CNN line later reaches 58/60 safe medium arrivals and 27/30 zero-shot hard arrivals, with 10/10 simple-turn retention. Its medium baseline is already 56/60, so the gain is two episodes, from one training seed. BEV supplies privileged scene structure and is not a camera encoder-pretraining result.')
    p('The retained best checkpoint has run counter 90k along a continued model lineage. Latest 150k collapses to 1/90 safe medium arrivals and 0/10 simple turns. The late policy mostly stalls. This failure shows why fewer collisions or more training steps alone cannot certify progress.')
    sub('C. Corrected label-efficiency comparison')
    fig('label_efficiency','Fig. 3. Corrected five-encoder validation score across actual label counts; mean over three seeds. Lower is better. Curves are offline development evidence, not route-success probabilities.')
    table('TABLE IV. Full-label policy metrics, three-seed means',
      ['Encoder','Score ↓','Stop recall','False holds'],[
      ['CNN scratch','0.3874','55.9%','18.2%'],
      ['ResNet-18','0.4477','36.8%','15.4%'],
      ['DINOv2-S','0.4323','44.9%','18.0%'],
      ['Comma JEPA','0.3755','52.0%','10.0%'],
      ['CARLA JEPA','0.4628','37.7%','13.6%'],
      ],[1.08,.62,.84,.84],8)
    p('Comma JEPA has the lowest full-label mean selector score, but its 5% nominal-budget false-hold rate is 51.1%. CARLA continuation does not improve the headline score here. The curves therefore support a qualified representation benefit, not a universal claim that in-domain JEPA is best or that high stop recall establishes competence.')

    # Page 5: native and planning.
    section(2)
    sub('D. Closed-loop pilots and a sustained camera drive')
    fig('native_pilots','Fig. 4. Full-label seed-42 camera pilots: arrivals versus strict successes out of six planned missions per encoder. Unknown cases remain in the denominator.')
    table('TABLE V. Native pilot outcomes; unfiltered contacts',
      ['Encoder','Arrived','Contacts','Red / stop','Unknown'],[
      ['CNN','0/6','1','1 / 0','2'],
      ['ResNet-18','6/6','0','1 / 1','0'],
      ['DINOv2-S','1/6','0','0 / 0','1'],
      ['CARLA JEPA','3/6','1','2 / 0','0'],
      ['Comma JEPA','4/6','10','2 / 1','0'],
      ],[1.0,.56,.60,.70,.52],7.8)
    p('All 30 main attempts are retained: 27 observed outcomes and three unknown failures. ResNet completes all six routes without observed contacts but violates one red light and one stop sign. Comma JEPA’s superior mean offline score does not produce the strongest native outcome. The original primary remains fixed rather than being reselected from this pilot.')
    p('Separately, a frozen ResNet camera policy completes a predetermined familiar Town04 TRAIN route of 1,375.64 m in 199.5 simulated seconds, with zero recorded contact starts and intervention ticks. Five minimum-speed infractions remain. Crucially, the route contains zero red-light and zero stop-sign opportunities; it demonstrates sustained camera navigation, not rule mastery. The paired expert crashes before its first trace, so an expert-relative comparison is unavailable.')
    p('The dedicated CNN-facts rule study records 15/72 passes, 22 contact starts, 15 red and five stop violations. None of 36 scenario pairs passes both arms with overrides disabled. These outcomes prevent a reliable-rule-understanding claim despite the successful familiar long drive.')
    col()
    sub('E. Action-conditioned world-model findings')
    p('The first-event study uses 68 independently verified recordings from 73 original attempts: 50 TRAIN episodes at eight locations and 18 calibration episodes at two locations. There are 12 TRAIN and four calibration contacts. Five failed attempts remain unknown. Future labels are censored at contact or intervention; scripted actors are nonreactive and all lights are green, limiting causal and rule coverage.')
    fig('world_model','Fig. 5. Calibration endpoint Brier score on the same 18 cases. The action-only baseline is strongest; learned JEPA and RSSM ensembles are not qualified for planning.')
    table('TABLE VI. Hazard prediction at the fixed 0.5 threshold',
      ['Model','Brier ↓','Missed contacts','False alarms'],[
      ['Constant','0.17285','4/4','0/14'],
      ['Action only','0.16328','4/4','0/14'],
      ['Calibrated CV','0.17597','4/4','0/14'],
      ['JEPA ensemble','0.17706','4/4','0/14'],
      ['RSSM ensemble','0.17728','4/4','0/14'],
      ],[1.14,.64,.86,.74],7.8)
    p('Six independent fits use 400 updates each. JEPA and RSSM ensemble members have about 454k and 623k parameters respectively, so matched updates do not imply equal compute. Both improve sequence likelihood relative to some simple controls but lose on endpoint Brier and miss every calibration contact. Lower likelihood alone does not justify planner deployment.')
    p('The raw CV proximity score flags all 18 cases: no missed contact, but 14 false alarms. Its calibrated version loses that sensitivity. Predictions for all fifteen candidate plans are retained, but outcomes for unexecuted options remain unknown. No counterfactual-ranking success is inferred from those missing labels, and no learned predictive planner is promoted.')

    # Page 6: perception and efficient implementation.
    section(2)
    sub('F. Full-distribution signal-perception audit')
    p('At low resolution, a distant bulb can occupy only a few pixels. The current signal pipeline therefore evaluates high-resolution point candidates with phase classification. A broader RGB proposal generator raises TRAIN candidate housing coverage to 1,246/1,248, but high proposal coverage alone does not solve foreground classification. Fig. 6 reports every fixed candidate, not only sampled training points.')
    fig('perception','Fig. 6. Full TRAIN signal tradeoff on the same 12.36 million candidates. Dashed lines mark minimum correct detections and maximum background detections. All three heads fail at least one required criterion.')
    p('The original plateau head yields 1,234 correct and 8,197 background detections. Adding 182,016 model-independent known-background examples reduces false detections to 4,848 but loses two red lights. Spatial context appends 4 × 4 RGB mean/standard-deviation grids at 16-, 32- and 96-pixel scales, increasing the descriptor from 87 to 375 features. This reduces false detections to 3,221, while correct detections fall to 1,228. The gate remains 1,234 and 3,075.8; neither regression is excused.')
    p('Independent recounts reproduce all town and phase metrics. Original 48,865 sampled feature/prediction rows and 182,016 added background rows match cached values exactly. The component validation has been repeatedly inspected during development and is not an untouched test set. Its separate best aggregate refinement of 55/72 remains rejected because yellow recall falls below the original protected value.')
    p('A subsequent TRAIN-only coverage audit identifies 1,596 frozen candidate coordinates that are exact audited active bulb pixels but absent from the old sampled supervision. All 4,788 original/gamma feature vectors pass the independent cache checks. This is released training data for a prospective comparison; no fit or improved model result is claimed in this paper snapshot.')
    col()
    sub('G. Exact preprocessing acceleration')
    p('Thousands of overlapping crop operations dominate the original RGB verifier. Precomputing aligned 3 × 3 integer averages reuses image work; preserving the original interleaved memory layout and sequential floating-point reduction order retains feature semantics. CPU compilation avoids materializing all patches. A CUDA implementation preserves sequential FP32 means and variances, with fused arithmetic disabled, then applies an exported numeric-tree head.')
    fig('runtime','Fig. 7. Three staged implementation profiles on the same 63 TRAIN images and 236,284 points. Includes image preparation and classifier work; CUDA includes transfers and returned probabilities. Excludes proposal generation and driving control.')
    p('All patch bytes and FP32/FP16 feature vectors match the original across the fixed 236,284-point audit. CPU probabilities are exact; maximum CUDA probability difference is 6.66 × 10⁻¹⁶, with no acceptance or phase changes. Median feature-plus-head time is 343.35 ms with shared averaging, 49.09 ms with compiled CPU reductions and 22.87 ms with CUDA plus GPU trees. CUDA p95 is 30.63 ms and peak allocated tensors are 11.09 MB.')
    p('The compiled CPU profile’s original control measures 972.83 ms median; the CUDA profile’s separate original control measures 1,638.27 ms. Different host conditions prevent a precise cross-run speedup claim. CUDA warm-up is 0.37 s. The retained concurrency annotation was corrected: its main audit had already finished before the CUDA profile, with only brief report/cache overlap.')
    p('These are staged kernel-path measurements. The newer 375-feature spatial head, three-camera scheduling, neural proposer, decoding, IPC, tracking and the policy require a separate end-to-end profile. Therefore, a 22.87 ms image path is not presented as a 10 Hz autonomous-driving system. Numerical equivalence and control reliability remain separate qualification questions.')
    sub('H. Outcome of the combined evidence')
    p('The project has a working simulator-control loop, useful frozen-encoder transfer, a sustained camera navigation demonstration and reproducible perception/planning audits. It does not yet have a qualified collision-free urban camera driver. Neither the 300-encounter camera gate nor the override-disabled rule suite has passed. The 36-mission Town05 final suite is registered but has not been run by a qualified final candidate.')

    # Page 7: discussion, next steps, references.
    section(2)
    h('VII. Research Gaps and Future Improvements')
    sub('A. Observability and decision-critical perception')
    p('Real highway pretraining does not cover every urban signal, stop sign, intersection layout or crossing interaction. Three RGB views expand observability, but they are not direct depth measurements or privileged 3D scene state. A detected bulb point supplies a calibrated ray, not a housing size, stop-line distance or route applicability. Future work must learn and verify lane-applicable signal association and temporal memory as the light leaves the camera view.')
    p('The immediate controlled data comparison is to add the 1,596 newly verified exact TRAIN positives while keeping negative examples, folds, class weights, model settings and gates fixed. If it succeeds on full TRAIN proposal distributions, the next steps are a registered component check, continuous approach–stop–release evaluation and only then policy retraining. A failed subgroup or background gate remains a failed result.')
    sub('B. Causal dynamics and calibrated planning')
    p('The action-conditioned models need more diverse observed consequences of alternative actions, including reactive crossing traffic, rear hazards and red-light outcomes. Prefix agreement must be measured before branch comparisons. Candidate rankings should be checked against executed alternative plans, with contact/intervention censoring preserved. JEPA should become a planning component only after it beats the declared constant-velocity and action-only baselines, including missed hazards rather than likelihood alone.')
    sub('C. Long-route reliability and benchmark discipline')
    p('Success compounds over encounters: if independent per-encounter safety is p, a route containing N such encounters succeeds with probability pᴺ. At p = 0.967 and N = 12, the illustrative route probability is about 0.67. Real encounters are not independent, so this is motivation for measuring encounter reliability, not a calibrated forecast of the camera driver. Short, well-covered tests and uninterrupted long drives are complementary.')
    p('The held-out Town05 suite contains twelve routes with three fixed seeds, spanning one to three kilometres. Planned signal/stop coverage and runtime opportunity checks must be retained. Late final-suite registration and prior Town03 engineering exposure are disclosed in project records; historical split-contaminated studies do not become clean tests by being renamed. Every collision remains in the headline result, and failed infrastructure cases are not replaced.')
    sub('D. Compute, stability and broader evidence')
    p('Next engineering priorities are full three-camera timing, compact atomic recordings, complete optimizer/scaler/RNG recovery state and diagnosis of the late Dreamer stopping collapse. Additional predeclared seeds and traffic streams are required before claiming a robust improvement. The privileged Dreamer BEV arm can remain a teacher/reference research line, but its performance must never be relabeled camera-only autonomy. Real-video open-loop checks could assess representation transfer without making a real-world driving claim.')
    col()
    h('VIII. Conclusion')
    p('Self-supervised video pretraining improves learning outcomes and weather robustness on the original constrained right-turn task, but those benefits do not automatically produce safe interactive driving. Corrected multi-camera studies reveal a gap between offline selection and native behavior, and learned hazard ensembles fail to beat simple baselines. Broader signal supervision and spatial context reduce false detections while exposing protected-recall regressions. Exact preprocessing makes substantially faster inference feasible without changing tested features or decisions. The central outcome is an auditable path from successful demonstrations to explicit unresolved requirements: reliable signal applicability, calibrated action consequences, rule-consistent control and held-out long-route evaluation. Final Town05 driving remains pending qualification.')
    p('Acknowledgment','Heading 5')
    p('[Funding, supervisor and institutional acknowledgment details; remove this paragraph if none.]',bold=True,size=9)
    p('References','Heading 5')
    refs=[
      'D. Hafner, J. Pasukonis, J. Ba, and T. Lillicrap, “Mastering diverse control tasks through world models,” Nature, vol. 640, pp. 647–653, 2025. doi: 10.1038/s41586-025-08744-2.',
      'A. Dosovitskiy, G. Ros, F. Codevilla, A. Lopez, and V. Koltun, “CARLA: An open urban driving simulator,” in Proc. Conf. Robot Learning, 2017, pp. 1–16. arXiv:1711.03938.',
      'D. Gao, S. Cai, H. Zhou, H. Wang, I. Soltani, and J. Zhang, “CarDreamer: Open-source learning platform for world model based autonomous driving,” arXiv:2405.09111, 2024.',
      'A. Bardes et al., “Revisiting feature prediction for learning visual representations from video,” Trans. Mach. Learn. Res., 2024. arXiv:2404.08471.',
      'M. Assran et al., “V-JEPA 2: Self-supervised video models enable understanding, prediction and planning,” arXiv:2506.09985, 2025.',
      'M. Oquab et al., “DINOv2: Learning robust visual features without supervision,” arXiv:2304.07193, 2023.',
      'K. He, X. Zhang, S. Ren, and J. Sun, “Deep residual learning for image recognition,” in Proc. IEEE Conf. Comput. Vis. Pattern Recognit., 2016, pp. 770–778. doi: 10.1109/CVPR.2016.90.',
      'Q. Li, X. Jia, S. Wang, and J. Yan, “Think2Drive: Efficient reinforcement learning by thinking in latent world model for quasi-realistic autonomous driving (in CARLA-v2),” in Proc. Eur. Conf. Comput. Vis., 2024. arXiv:2402.16720.',
      'X. Jia, Z. Yang, Q. Li, Z. Zhang, and J. Yan, “Bench2Drive: Towards multi-ability benchmarking of closed-loop end-to-end autonomous driving,” Adv. Neural Inf. Process. Syst., Datasets and Benchmarks Track, 2024. arXiv:2406.03877.',
      'comma.ai, “comma2k19,” dataset and documentation, 2018. [Online]. Available: https://github.com/commaai/comma2k19. Accessed: Oct. 23, 2026.',
      'Auto Thinker Engine project, “Versioned experimental contracts, corrected encoder study, native audits and journey log,” unpublished project artifacts, Oct. 2026. Evidence snapshot: Oct. 23, 2026. Local index: reports/navigation_results_index_20261020.md.',
    ]
    for i,text in enumerate(refs,1):
        a=p(f'[{i}] {text}','references');a.paragraph_format.left_indent=Pt(11)
        a.paragraph_format.first_line_indent=Pt(-11);a.paragraph_format.space_after=Pt(4)
    DOC.save(DEST)
    manifest={str(p):hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in SOURCES}
    (OUT/'source_manifest.json').write_text(json.dumps({'snapshot':'2026-10-23',
      'template':'/home/deepakbalaji/Downloads/conference-template-a4.docx',
      'template_sha256':hashlib.sha256(Path('/home/deepakbalaji/Downloads/conference-template-a4.docx').read_bytes()).hexdigest(),
      'sources':manifest,'only_unfilled_fields_bold':True,'new_training_started':False,
      'references':['https://www.nature.com/articles/s41586-025-08744-2',
       'https://arxiv.org/abs/1711.03938','https://arxiv.org/abs/2405.09111',
       'https://arxiv.org/abs/2404.08471','https://arxiv.org/abs/2506.09985',
       'https://arxiv.org/abs/2304.07193','https://arxiv.org/abs/1512.03385',
       'https://arxiv.org/abs/2402.16720','https://arxiv.org/abs/2406.03877',
       'https://github.com/commaai/comma2k19'],
      'final_docx':str(DEST.relative_to(ROOT))},indent=2)+'\n')
    (OUT/'paper_text.txt').write_text('\n\n'.join(a.text for a in DOC.paragraphs))
    print('DOCX:',DEST,'bytes:',DEST.stat().st_size)


if __name__=='__main__':
    figures()
    build()
